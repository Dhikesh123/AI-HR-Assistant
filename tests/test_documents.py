"""Document upload, validation, re-index and delete tests."""
from __future__ import annotations

import io
from pathlib import Path

from fastapi.testclient import TestClient

SAMPLE_DIR = Path(__file__).resolve().parents[1] / "data" / "sample_hr_documents"


def _upload(client: TestClient, headers: dict, name: str, content: bytes, mime: str):
    return client.post(
        "/api/documents/upload",
        headers=headers,
        files={"file": (name, io.BytesIO(content), mime)},
    )


def test_upload_requires_admin(client: TestClient, employee_headers: dict) -> None:
    response = _upload(client, employee_headers, "policy.txt", b"Some policy text.", "text/plain")
    assert response.status_code == 403


def test_upload_pdf(client: TestClient, admin_headers: dict) -> None:
    content = (SAMPLE_DIR / "holiday_policy.pdf").read_bytes()
    response = _upload(client, admin_headers, "holiday_policy.pdf", content, "application/pdf")
    assert response.status_code == 201, response.text
    document = response.json()["document"]
    assert document["status"] == "indexed"
    assert document["chunk_count"] > 0
    assert document["file_type"] == ".pdf"


def test_upload_docx(client: TestClient, admin_headers: dict, tmp_path: Path) -> None:
    from docx import Document as DocxDocument

    path = tmp_path / "travel_policy.docx"
    doc = DocxDocument()
    doc.add_heading("1. Travel Allowance", level=1)
    doc.add_paragraph(
        "Employees travelling on company business receive a daily allowance of INR 2,500."
    )
    doc.save(path)

    response = _upload(
        client,
        admin_headers,
        "travel_policy.docx",
        path.read_bytes(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert response.status_code == 201, response.text
    assert response.json()["document"]["status"] == "indexed"


def test_upload_txt_becomes_answerable(client: TestClient, admin_headers: dict, employee_headers: dict) -> None:
    content = (
        "1. Sabbatical Policy\n"
        "Employees who complete five years of service may apply for a sabbatical "
        "of up to 90 days without pay.\n"
    ).encode()
    upload = _upload(client, admin_headers, "sabbatical_policy.txt", content, "text/plain")
    assert upload.status_code == 201

    answer = client.post(
        "/api/chat",
        headers=employee_headers,
        json={"question": "Can employees apply for a sabbatical without pay?"},
    ).json()
    assert "sabbatical_policy.txt" in {source["document"] for source in answer["sources"]}


def test_upload_rejects_invalid_file_type(client: TestClient, admin_headers: dict) -> None:
    response = _upload(client, admin_headers, "malware.exe", b"MZ binary", "application/x-msdownload")
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["detail"]


def test_upload_rejects_empty_file(client: TestClient, admin_headers: dict) -> None:
    response = _upload(client, admin_headers, "empty.txt", b"", "text/plain")
    assert response.status_code == 400


def test_upload_rejects_oversized_file(client: TestClient, admin_headers: dict) -> None:
    from backend.core.config import settings

    oversized = b"x" * (settings.max_upload_bytes + 1024)
    response = _upload(client, admin_headers, "huge.txt", oversized, "text/plain")
    assert response.status_code == 400
    assert "too large" in response.json()["detail"].lower()


def test_list_and_get_documents(client: TestClient, admin_headers: dict) -> None:
    _upload(client, admin_headers, "listed_policy.txt", b"1. Parking\nParking is free.", "text/plain")

    listing = client.get("/api/documents", headers=admin_headers)
    assert listing.status_code == 200
    documents = listing.json()
    assert documents

    document_id = documents[0]["id"]
    detail = client.get(f"/api/documents/{document_id}", headers=admin_headers)
    assert detail.status_code == 200
    assert detail.json()["id"] == document_id


def test_get_missing_document_returns_404(client: TestClient, admin_headers: dict) -> None:
    assert client.get("/api/documents/999999", headers=admin_headers).status_code == 404


def test_reindex_document(client: TestClient, admin_headers: dict) -> None:
    upload = _upload(
        client,
        admin_headers,
        "reindex_policy.txt",
        b"1. Gym Access\nThe office gym is open from 6 AM to 10 PM on weekdays.",
        "text/plain",
    )
    document_id = upload.json()["document"]["id"]
    chunk_count = upload.json()["document"]["chunk_count"]

    response = client.post(f"/api/documents/{document_id}/reindex", headers=admin_headers)
    assert response.status_code == 200
    assert response.json()["status"] == "indexed"
    # Re-indexing replaces vectors rather than duplicating them.
    assert response.json()["chunk_count"] == chunk_count


def test_delete_document_removes_it_from_retrieval(
    client: TestClient, admin_headers: dict, employee_headers: dict
) -> None:
    upload = _upload(
        client,
        admin_headers,
        "shuttle_policy.txt",
        b"1. Shuttle Service\nA free shuttle runs between the metro station and the office.",
        "text/plain",
    )
    document_id = upload.json()["document"]["id"]

    before = client.post(
        "/api/chat", headers=employee_headers, json={"question": "Is there a free shuttle service?"}
    ).json()
    assert "shuttle_policy.txt" in {source["document"] for source in before["sources"]}

    assert client.delete(f"/api/documents/{document_id}", headers=admin_headers).status_code == 204
    assert client.get(f"/api/documents/{document_id}", headers=admin_headers).status_code == 404

    after = client.post(
        "/api/chat", headers=employee_headers, json={"question": "Is there a free shuttle service?"}
    ).json()
    assert "shuttle_policy.txt" not in {source["document"] for source in after["sources"]}


def test_delete_requires_admin(client: TestClient, employee_headers: dict) -> None:
    assert client.delete("/api/documents/1", headers=employee_headers).status_code == 403
