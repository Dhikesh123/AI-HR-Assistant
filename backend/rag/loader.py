"""Document loading and text extraction.

Supports PDF, DOCX, TXT and Markdown. Each loader returns a list of
``(page_number, text)`` tuples so page-level citations stay possible.
"""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path
from typing import List, Tuple

from backend.core.logging_config import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

Page = Tuple[int, str]

# A wrapped line is long; a short line is usually a heading or a list item and
# must keep its own line so section detection can see it.
_WRAP_THRESHOLD = 55


def clean_text(text: str) -> str:
    """Normalise whitespace and rejoin lines that PDF extraction hard-wrapped.

    Short lines are preserved as-is because they are usually headings or list
    items, which the chunker relies on for section metadata.
    """
    text = text.replace("\x00", " ").replace("﻿", "")
    text = re.sub(r"[ \t]+", " ", text)

    lines = [line.strip() for line in text.splitlines()]
    output: List[str] = []
    buffer = ""

    for line in lines:
        if not line:
            if buffer:
                output.append(buffer)
                buffer = ""
            output.append("")
            continue

        if not buffer:
            buffer = line
            continue

        # A long line that does not end a sentence was wrapped by the PDF
        # renderer, so it continues into the next line. A short line is a
        # heading or list item and keeps its own line.
        wrapped = len(buffer) >= _WRAP_THRESHOLD and not buffer.endswith(
            (".", ":", ";", "!", "?")
        )
        if wrapped:
            buffer = f"{buffer} {line}"
        else:
            output.append(buffer)
            buffer = line

    if buffer:
        output.append(buffer)

    result = "\n".join(output)
    return re.sub(r"\n{3,}", "\n\n", result).strip()


def strip_repeated_lines(pages: List[Page]) -> List[Page]:
    """Remove running headers/footers that repeat across most pages.

    Boilerplate like a footer on every page adds no information but does dilute
    embeddings and can be mistaken for a section heading.
    """
    if len(pages) < 3:
        return pages

    counts: Counter[str] = Counter()
    for _, text in pages:
        for line in {line.strip() for line in text.splitlines() if line.strip()}:
            counts[line] += 1

    threshold = max(2, int(len(pages) * 0.6))
    boilerplate = {line for line, count in counts.items() if count >= threshold and len(line) < 120}
    if not boilerplate:
        return pages

    cleaned: List[Page] = []
    for page_number, text in pages:
        kept = "\n".join(l for l in text.splitlines() if l.strip() not in boilerplate)
        kept = re.sub(r"\n{3,}", "\n\n", kept).strip()
        if kept:
            cleaned.append((page_number, kept))
    return cleaned or pages


class UnsupportedFileTypeError(ValueError):
    """Raised when a file extension is not supported."""


class DocumentExtractionError(RuntimeError):
    """Raised when text could not be extracted from a document."""


def _load_pdf(path: Path) -> List[Page]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages: List[Page] = []
    for index, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception as exc:  # pragma: no cover - defensive
            logger.warning("Failed to extract page %s of %s: %s", index, path.name, exc)
            raw = ""
        cleaned = clean_text(raw)
        if cleaned:
            pages.append((index, cleaned))
    return strip_repeated_lines(pages)


def _load_docx(path: Path) -> List[Page]:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    blocks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))
    text = clean_text("\n".join(blocks))
    return [(1, text)] if text else []


def _load_text(path: Path) -> List[Page]:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    text = clean_text(raw)
    return [(1, text)] if text else []


def load_document(file_path: str | Path) -> List[Page]:
    """Extract text from ``file_path`` as a list of ``(page_number, text)``."""
    path = Path(file_path)
    if not path.exists():
        raise DocumentExtractionError(f"File not found: {path}")

    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(f"Unsupported file type: {suffix}")

    try:
        if suffix == ".pdf":
            pages = _load_pdf(path)
        elif suffix == ".docx":
            pages = _load_docx(path)
        else:
            pages = _load_text(path)
    except (UnsupportedFileTypeError, DocumentExtractionError):
        raise
    except Exception as exc:
        raise DocumentExtractionError(f"Could not extract text from {path.name}: {exc}") from exc

    if not pages:
        raise DocumentExtractionError(f"No readable text found in {path.name}")

    logger.info("Extracted %s page(s) of text from %s", len(pages), path.name)
    return pages
